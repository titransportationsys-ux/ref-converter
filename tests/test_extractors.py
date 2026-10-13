import shutil
import zipfile
from pathlib import Path

import pytest

import fixtures as F
from conftest import codes
from refconvert import extractors, pipeline
from refconvert.model import ExtractError, Unsupported

PAGES = [{"lines": ["Chapter 3 Design speed", "The minimum radius for 80 km/h is 230 m.", "Superelevation e max = 0.08"],
          "table": [["Speed", "Radius"], ["60", "120"], ["80", "230"]]},
         {"lines": ["Page two has 1,250 veh/h and 3.5 % grade."]}]


def test_pdf_engines_agree_and_clean_table_is_corroborated_by_line_layout(make):
    _, w = make("pdf", pages=PAGES)
    assert len(w.merged()["engines"]) >= 2
    assert [u["unit"] for u in w.merged()["units"]] == ["p1", "p2"]
    assert codes(w, "open", "BLOCK") == []
    tb = next(b for b in w.merged()["units"][0]["blocks"] if b["kind"] == "table")
    assert tb["rows"] == [["Speed", "Radius"], ["60", "120"], ["80", "230"]]
    assert any(a.endswith(":line-layout") for a in tb["meta"]["agreed"])


def test_pdf_table_whose_layout_cannot_be_corroborated_blocks(make):
    pages = [dict(PAGES[0], skew=True)]
    _, w = make("pdf", pages=pages)
    assert codes(w, "open", "BLOCK") == ["TABLE_STRUCTURE"]          # numbers are on the page, alignment is not proven
    assert w.merged()["units"][0]["status"] == "unverified"


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx", "csv", "html", "bib", "ris", "dxf", "md", "json"])
def test_clean_formats_have_no_blocking_flags(make, kind):
    _, w = make(kind)
    assert codes(w, "open", "BLOCK") == [], w.review()["flags"]
    assert all(u["status"] == "verified" for u in w.merged()["units"])
    if kind not in ("md", "json"):
        assert len(w.merged()["engines"]) >= 2


def test_docx_content_is_complete(make):
    _, w = make("docx")
    text = " ".join(b["text"] for u in w.merged()["units"] for b in u["blocks"] if b["kind"] == "text")
    assert "230 m" in text and "0.278Vt" in text
    tables = [b for u in w.merged()["units"] for b in u["blocks"] if b["kind"] == "table"]
    assert tables[0]["rows"] == [["Speed", "Radius", "e max"], ["60", "120", "0.08"], ["80", "230", "0.08"]]


def test_xlsx_flags_hidden_sheet_dates_and_formula_without_cache(make):
    _, w = make("xlsx")
    assert {"HIDDEN_SHEET", "DATE_CELLS"} <= set(codes(w))
    _, w2 = make("xlsx", name="nocache.xlsx", formula_nocache=True)
    assert "FORMULA_NO_CACHE" in codes(w2, "open", "BLOCK")


def test_csv_quoted_comma_survives(make):
    _, w = make("csv")
    rows = w.merged()["units"][0]["blocks"][0]["rows"]
    assert rows[1] == ["Arterial", "80", "fast, wide"]


def test_json_preserves_number_text(make):
    _, w = make("json")
    rows = w.merged()["units"][0]["blocks"][0]["rows"]
    assert ["$.links[0].speed", "80.50"] in rows       # trailing zero kept


def test_bib_and_ris_fields(make):
    _, w = make("bib")
    assert [u["unit"] for u in w.merged()["units"]] == ["bib:basyoni2016", "bib:aashto2018"]
    _, w2 = make("ris")
    assert len(w2.merged()["units"]) == 2


def test_dxf_entities(make):
    _, w = make("dxf")
    rows = w.merged()["units"][0]["blocks"][0]["rows"]
    types = [r[0] for r in rows[1:]]
    assert types == ["LINE", "CIRCLE", "TEXT", "LWPOLYLINE"]
    assert ["TEXT", "ANNO", "R=25 m", "10 10 0"] in rows


def test_image_needs_transcription_then_blind_pair_resolves(make, tmp_path):
    _, w = make("png")
    assert codes(w, "open", "BLOCK") == ["NEEDS_TRANSCRIPTION"]
    a = tmp_path / "a.yaml"
    a.write_text("units:\n  'img:src': 'Design speed 80 km/h, radius 230 m'\n")
    b = tmp_path / "b.yaml"
    b.write_text("units:\n  'img:src': 'Design speed 80 km/h, radius 230 m'\n")
    pipeline.ingest(w, "vision-A", a)
    assert codes(w, "open", "BLOCK") == ["SINGLE_SOURCE"]                 # one transcription is not enough
    pipeline.ingest(w, "vision-B", b)
    assert codes(w, "open", "BLOCK") == []
    assert w.merged()["units"][0]["status"] == "verified"


def test_blind_transcriptions_that_disagree_block(make, tmp_path):
    _, w = make("png")
    for eng, txt in (("vision-A", "radius 230 m"), ("vision-B", "radius 280 m")):
        f = tmp_path / f"{eng}.yaml"
        f.write_text(f"units:\n  'img:src': '{txt}'\n")
        pipeline.ingest(w, eng, f)
    assert codes(w, "open", "BLOCK") == ["NO_CONSENSUS"]
    assert "230" in w.review()["flags"][-1]["detail"] or any("230" in f["detail"] for f in w.review()["flags"])


def test_empty_drawing_page_blocks(make):
    _, w = make("pdf", pages=[{"lines": [], "graphics": 40}])
    assert codes(w, "open", "BLOCK") == ["EMPTY_UNIT_GRAPHICS"]


def test_docx_tracked_change_makes_engines_disagree(make, tmp_path):
    src = F.make_docx(tmp_path / "tc.docx")
    out = tmp_path / "tc2.docx"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(out, "w") as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b"</w:p>", b'<w:ins w:id="9" w:author="x" w:date="2024-01-01T00:00:00Z"><w:r><w:t> 777</w:t></w:r></w:ins></w:p>', 1)
            zo.writestr(item, data)
    w = pipeline.extract(out, tmp_path / "work_tc")
    assert "TRACKED_CHANGES" in codes(w)
    assert "NO_CONSENSUS" in codes(w, "open", "BLOCK")


def test_unsupported_and_spoofed_files(tmp_path):
    (tmp_path / "x.dwg").write_bytes(b"AC1027")
    with pytest.raises(Unsupported, match="DXF"):
        extractors.detect(tmp_path / "x.dwg")
    (tmp_path / "fake.pdf").write_bytes(b"not a pdf")
    with pytest.raises(ExtractError, match="not a PDF"):
        extractors.detect(tmp_path / "fake.pdf")
    (tmp_path / "blob.bin").write_bytes(b"\x00\x01\x02")
    with pytest.raises(Unsupported):
        extractors.detect(tmp_path / "blob.bin")


def test_existing_work_folder_is_never_overwritten(make, tmp_path):
    src, w = make("md")
    with pytest.raises(ExtractError, match="already holds"):
        pipeline.extract(src, w.root)


def test_legacy_doc_converts_when_libreoffice_present(tmp_path):
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed")
    d = F.make_docx(tmp_path / "s.docx")
    import subprocess
    subprocess.run(["soffice", "--headless", "--convert-to", "doc", "--outdir", str(tmp_path), str(d)], capture_output=True, timeout=180)
    legacy = tmp_path / "s.doc"
    assert legacy.exists()
    w = pipeline.extract(legacy, tmp_path / "work_legacy")
    assert w.manifest()["derived_from"].endswith("s.doc")


def test_merged_cells_keep_columns_aligned(tmp_path):
    import docx
    d = docx.Document()
    t = d.add_table(rows=2, cols=3)
    for r in range(2):
        for c in range(3):
            t.cell(r, c).text = f"{r}{c}"
    t.cell(0, 0).merge(t.cell(0, 1))
    d.save(str(tmp_path / "m.docx"))
    w = pipeline.extract(tmp_path / "m.docx", tmp_path / "w")
    rows = next(b for b in w.merged()["units"][0]["blocks"] if b["kind"] == "table")["rows"]
    assert rows[0] == ["00\n01", "", "02"] and len(rows[1]) == 3        # '02' stays in column 3
    h = tmp_path / "t.html"
    h.write_text('<table><tr><th colspan="2">Speed</th><th>R</th></tr><tr><td>1</td><td>2</td><td>3</td></tr></table>')
    w2 = pipeline.extract(h, tmp_path / "wh")
    assert next(b for b in w2.merged()["units"][0]["blocks"] if b["kind"] == "table")["rows"][0] == ["Speed", "", "R"]


def test_ragged_table_is_warned(tmp_path):
    src = tmp_path / "r.csv"
    src.write_text("a,b,c\n1,2\n")
    w = pipeline.extract(src, tmp_path / "w")
    assert codes(w) == ["DELIMITER"] or "TABLE_RAGGED" not in codes(w)   # csv pads rows to a rectangle, so no ragged table
    h = tmp_path / "t.html"
    h.write_text("<table><tr><td>1</td><td>2</td><td>3</td></tr><tr><td rowspan='2'>x</td><td>y</td></tr><tr><td>z</td></tr></table>")
    w2 = pipeline.extract(h, tmp_path / "w2")
    assert "TABLE_RAGGED" in codes(w2, "open", "WARN")


def test_html_inline_markup_and_table_cells_do_not_split_or_glue_numbers(tmp_path):
    h = tmp_path / "t.html"
    h.write_text("<html><body><p>Flow <b>1</b>,250 veh/h &amp; <i>3</i>.5 % grade</p><!-- 999 -->"
                 "<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>")
    w = pipeline.extract(h, tmp_path / "w")
    assert codes(w, "open", "BLOCK") == []
    text = " ".join(b["text"] for b in w.merged()["units"][0]["blocks"] if b["kind"] == "text")
    assert "1,250 veh/h & 3.5 % grade" in text and "999" not in text
