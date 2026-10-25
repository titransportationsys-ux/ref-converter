"""Build small real files in every supported format for the tests."""
from __future__ import annotations

from pathlib import Path


def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages: list[dict]) -> Path:
    """pages: [{"lines": [...], "table": [[...]]?, "graphics": int?}] -> minimal valid PDF with text, ruled tables."""
    n = len(pages)
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            ("<< /Type /Pages /Kids [" + " ".join(f"{6 + 2 * i} 0 R" for i in range(n)) + f"] /Count {n} >>").encode(),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Wingdings /FirstChar 32 /LastChar 126 /Widths [" + b" ".join([b"800"] * 95)
            + b"] /FontDescriptor 5 0 R >>",
            b"<< /Type /FontDescriptor /FontName /Wingdings /Flags 4 /FontBBox [0 0 1000 1000] /ItalicAngle 0 /Ascent 900 /Descent -200 "
            b"/CapHeight 700 /StemV 80 >>"]
    for i, p in enumerate(pages):
        s = []
        if p.get("lines"):
            s.append("BT /F1 11 Tf 50 780 Td 15 TL")
            for ln in p["lines"]:
                s.append(f"({_esc(ln)}) Tj T*")
            s.append("ET")
        t = p.get("table")
        if t:
            x0, y0, cw, ch = 50, 560, 110, 22
            rows, cols = len(t), len(t[0])
            for r in range(rows + 1):
                s.append(f"{x0} {y0 - r * ch} m {x0 + cols * cw} {y0 - r * ch} l S")
            for c in range(cols + 1):
                s.append(f"{x0 + c * cw} {y0} m {x0 + c * cw} {y0 - rows * ch} l S")
            for r, row in enumerate(t):
                for c, cell in enumerate(row):
                    sup = None
                    if isinstance(cell, tuple):                      # (text, superscript)
                        cell, sup = cell
                    drop = 9 if (p.get("skew") and r == 1 and c == 0) else 0       # one cell sits lower: breaks line layout
                    s.append(f"BT /F1 10 Tf {x0 + c * cw + 4} {y0 - r * ch - 15 + 8 - drop} Td ({_esc(cell)}) Tj ET")
                    if sup:
                        s.append(f"BT /F1 6 Tf {x0 + c * cw + 4 + 5.56 * len(cell)} {y0 - r * ch - 15 + 8 - drop + 4} Td ({_esc(sup)}) Tj ET")
        if p.get("raw"):
            s.append(p["raw"])
        for g in range(p.get("graphics", 0)):
            s.append(f"{60 + g * 5} 100 m {60 + g * 5} 150 l S")
        stream = "\n".join(s).encode("latin-1")
        objs.append((f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> "
                     f"/Contents {7 + 2 * i} 0 R >>").encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for k, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{k} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def make_docx(path: Path) -> Path:
    import docx
    d = docx.Document()
    d.add_heading("Design speed", level=1)
    d.add_paragraph("The minimum radius for 80 km/h is 230 m (Table 3-1).")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["Speed", "Radius", "e max"], ["60", "120", "0.08"], ["80", "230", "0.08"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    t.cell(0, 0).merge(t.cell(0, 1)) if False else None
    d.add_paragraph("Stopping sight distance SSD = 0.278Vt + V^2/(254(a/9.81 + G))")
    d.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Lane width 3.65 m"
    tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    tb.text_frame.text = "Shoulder 2.5 m on 80 km/h roads"
    gt = s.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for r, row in enumerate([["Class", "Width"], ["A", "3.65"]]):
        for c, v in enumerate(row):
            gt.cell(r, c).text = v
    p.save(str(path))
    return path


def make_xlsx(path: Path, formula_nocache=False) -> Path:
    import datetime
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Speeds"
    ws.append(["Class", "Speed", "Note"])
    ws.append(["Arterial", 80, "ok"])
    ws.append(["Local", 40.5, ""])
    ws["D1"] = "when"
    ws["D2"] = datetime.datetime(2024, 3, 1)
    if formula_nocache:
        ws["B4"] = "=B2*2"
    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "secret 12"
    hidden.sheet_state = "hidden"
    wb.save(str(path))
    return path


def make_csv(path: Path) -> Path:
    path.write_text("Class,Speed,Note\nArterial,80,\"fast, wide\"\nLocal,40.5,\n", encoding="utf-8")
    return path


def make_html(path: Path) -> Path:
    path.write_text("<html><head><title>x</title><script>var a=99;</script></head><body><h1>Capacity</h1>"
                    "<p>Saturation flow is 1900 pcu/h.</p><table><tr><th>Lane</th><th>Flow</th></tr>"
                    "<tr><td>Shared</td><td>1750</td></tr></table></body></html>", encoding="utf-8")
    return path


def make_bib(path: Path) -> Path:
    path.write_text('@article{basyoni2016,\n  author = {Basyoni, A. and Talaat, H.},\n  title = {Traffic states {B}ayesian data expansion},\n'
                    '  year = 2016,\n  journal = "ITS Congress",\n  pages = {1--10}\n}\n@book{aashto2018, title={A Policy on Geometric Design}, '
                    'publisher={AASHTO}, year={2018}, edition={7}}\n', encoding="utf-8")
    return path


def make_ris(path: Path) -> Path:
    path.write_text("TY  - JOUR\nAU  - Basyoni, A.\nAU  - Talaat, H.\nTI  - Traffic states data expansion\nPY  - 2016\nSP  - 1\nEP  - 10\nER  - \n"
                    "TY  - BOOK\nTI  - Green Book\nPY  - 2018\nER  - \n", encoding="utf-8")
    return path


def make_dxf(path: Path) -> Path:
    import ezdxf
    doc = ezdxf.new("R2010")
    msp = doc.modelspace()
    msp.add_line((0, 0), (100.5, 0), dxfattribs={"layer": "CL"})
    msp.add_circle((50, 50), 25, dxfattribs={"layer": "RAB"})
    msp.add_text("R=25 m", dxfattribs={"layer": "ANNO", "insert": (10, 10), "height": 2})
    msp.add_lwpolyline([(0, 0), (10, 0), (10, 10)], close=True, dxfattribs={"layer": "EDGE"})
    doc.saveas(str(path))
    return path


def make_text(path: Path) -> Path:
    path.write_text("# Rules\n\nThe design speed shall be 80 km/h.\nMinimum radius = 230 m.\n", encoding="utf-8")
    return path


def make_json(path: Path) -> Path:
    path.write_text('{"links": [{"id": 1, "speed": 80.50, "up": [13]}, {"id": 2, "speed": 40}], "n": 2}', encoding="utf-8")
    return path


def make_png(path: Path) -> Path:
    from PIL import Image
    Image.new("RGB", (20, 10), "white").save(path)
    return path
